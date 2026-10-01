import os
from typing import Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_to_excel(data: Dict[str, Any], filepath: str) -> str:
    """
    Exports scraped EGX investor data to a beautifully formatted Excel file.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styles
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="0A3D62", end_color="0A3D62", fill_type="solid")
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3799", end_color="1E3799", fill_type="solid")
    
    meta_font = Font(name="Calibri", size=10, italic=True, color="333333")
    
    green_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    green_font = Font(name="Calibri", size=11, bold=True, color="1B5E20")
    
    red_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
    red_font = Font(name="Calibri", size=11, bold=True, color="B71C1C")
    
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    sheets_config = [
        ("الإجمالي العام", "all", "تقرير إجمالي تعاملات المستثمرين (الكل)"),
        ("الأفراد", "individuals", "تقرير تعاملات المستثمرين الأفراد"),
        ("المؤسسات", "institutions", "تقرير تعاملات المستثمرين المؤسسات")
    ]

    for sheet_title, key, full_title in sheets_config:
        ws = wb.create_sheet(title=sheet_title)
        ws.sheet_view.rightToLeft = True

        # Header Title
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = f"البورصة المصرية - {full_title}"
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = align_center
        ws.row_dimensions[1].height = 32

        # Metadata
        ws.merge_cells("A2:G2")
        meta_cell = ws["A2"]
        meta_cell.value = f"القطاع: {data.get('segment_name')} | تاريخ ووقت السحب: {data.get('timestamp')} | سعر الصرف المعتمد: 1 USD = {data.get('usd_rate', 52):.2f} EGP"
        meta_cell.font = meta_font
        meta_cell.alignment = align_center
        ws.row_dimensions[2].height = 22

        # Table Column Headers
        headers = [
            "فئة المستثمر",
            "قيمة الشراء (EGP)",
            "قيمة البيع (EGP)",
            "صافي القيمة (EGP)",
            "قيمة الشراء (USD)",
            "قيمة البيع (USD)",
            "صافي القيمة (USD)"
        ]

        ws.row_dimensions[4].height = 26
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border

        # Rows
        items = data.get("tables", {}).get(key, [])
        for row_idx, item in enumerate(items, 5):
            ws.row_dimensions[row_idx].height = 24
            
            c_type = ws.cell(row=row_idx, column=1, value=item["type"])
            c_type.alignment = align_center
            c_type.font = Font(name="Calibri", size=11, bold=True)
            c_type.border = thin_border

            # EGP values
            c_buy_egp = ws.cell(row=row_idx, column=2, value=item["buy_egp"])
            c_buy_egp.number_format = '#,##0'
            c_buy_egp.alignment = align_right
            c_buy_egp.border = thin_border

            c_sell_egp = ws.cell(row=row_idx, column=3, value=item["sell_egp"])
            c_sell_egp.number_format = '#,##0'
            c_sell_egp.alignment = align_right
            c_sell_egp.border = thin_border

            c_net_egp = ws.cell(row=row_idx, column=4, value=item["net_egp"])
            c_net_egp.number_format = '#,##0'
            c_net_egp.alignment = align_right
            c_net_egp.border = thin_border

            # USD values
            c_buy_usd = ws.cell(row=row_idx, column=5, value=item["buy_usd"])
            c_buy_usd.number_format = '$#,##0.00'
            c_buy_usd.alignment = align_right
            c_buy_usd.border = thin_border

            c_sell_usd = ws.cell(row=row_idx, column=6, value=item["sell_usd"])
            c_sell_usd.number_format = '$#,##0.00'
            c_sell_usd.alignment = align_right
            c_sell_usd.border = thin_border

            c_net_usd = ws.cell(row=row_idx, column=7, value=item["net_usd"])
            c_net_usd.number_format = '$#,##0.00'
            c_net_usd.alignment = align_right
            c_net_usd.border = thin_border

            # Conditional colors for net
            if item["net_egp"] >= 0:
                c_net_egp.fill = green_fill
                c_net_egp.font = green_font
                c_net_usd.fill = green_fill
                c_net_usd.font = green_font
            else:
                c_net_egp.fill = red_fill
                c_net_egp.font = red_font
                c_net_usd.fill = red_fill
                c_net_usd.font = red_font

        # Auto-adjust column width
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if cell.row > 2 and len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
    wb.save(filepath)
    return filepath

def export_archive_to_excel(sb: str, filepath: str) -> str:
    """
    Exports daily, weekly, and monthly institutional flow archive to Excel.
    """
    from core.archive import get_daily_history, get_weekly_summary, get_monthly_summary

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Styles
    title_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="0A3D62", end_color="0A3D62", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3799", end_color="1E3799", fill_type="solid")
    green_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    green_font = Font(name="Calibri", size=10, bold=True, color="1B5E20")
    red_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
    red_font = Font(name="Calibri", size=10, bold=True, color="B71C1C")
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD')
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. Daily Sheet
    ws_daily = wb.create_sheet(title="المعاملات اليومية للمؤسسات")
    ws_daily.sheet_view.rightToLeft = True

    ws_daily.merge_cells("A1:H1")
    ws_daily["A1"].value = "أرشيف معاملات المؤسسات اليومي (دخل / خرج / صافي)"
    ws_daily["A1"].font = title_font
    ws_daily["A1"].fill = title_fill
    ws_daily["A1"].alignment = align_center

    daily_headers = [
        "التاريخ", "دخل (شراء EGP)", "خرج (بيع EGP)", "صافي التدفق (EGP)",
        "دخل (USD)", "خرج (USD)", "صافي التدفق (USD)", "صافي الأجانب (EGP)"
    ]
    for col_idx, h in enumerate(daily_headers, 1):
        c = ws_daily.cell(row=3, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_center
        c.border = thin_border

    daily_rows = get_daily_history(sb, limit=100)
    for r_idx, r in enumerate(daily_rows, 4):
        ws_daily.cell(row=r_idx, column=1, value=r["date"]).alignment = align_center
        
        c_in_egp = ws_daily.cell(row=r_idx, column=2, value=r["total_inflow_egp"])
        c_in_egp.number_format = '#,##0'
        c_out_egp = ws_daily.cell(row=r_idx, column=3, value=r["total_outflow_egp"])
        c_out_egp.number_format = '#,##0'
        
        c_net_egp = ws_daily.cell(row=r_idx, column=4, value=r["total_net_egp"])
        c_net_egp.number_format = '#,##0'
        c_net_egp.fill = green_fill if r["total_net_egp"] >= 0 else red_fill
        c_net_egp.font = green_font if r["total_net_egp"] >= 0 else red_font

        c_in_usd = ws_daily.cell(row=r_idx, column=5, value=r["total_inflow_usd"])
        c_in_usd.number_format = '$#,##0.00'
        c_out_usd = ws_daily.cell(row=r_idx, column=6, value=r["total_outflow_usd"])
        c_out_usd.number_format = '$#,##0.00'

        c_net_usd = ws_daily.cell(row=r_idx, column=7, value=r["total_net_usd"])
        c_net_usd.number_format = '$#,##0.00'
        c_net_usd.fill = green_fill if r["total_net_usd"] >= 0 else red_fill
        c_net_usd.font = green_font if r["total_net_usd"] >= 0 else red_font

        c_f_egp = ws_daily.cell(row=r_idx, column=8, value=r["foreign_net_egp"])
        c_f_egp.number_format = '#,##0'
        c_f_egp.fill = green_fill if r["foreign_net_egp"] >= 0 else red_fill
        c_f_egp.font = green_font if r["foreign_net_egp"] >= 0 else red_font

    # 2. Weekly Sheet
    ws_week = wb.create_sheet(title="المقارنة الأسبوعية")
    ws_week.sheet_view.rightToLeft = True
    ws_week.merge_cells("A1:G1")
    ws_week["A1"].value = "ملخص تدفقات المؤسسات أسبوعياً"
    ws_week["A1"].font = title_font
    ws_week["A1"].fill = title_fill
    ws_week["A1"].alignment = align_center

    weekly_headers = ["الأسبوع", "أيام التداول", "إجمالي دخل (EGP)", "إجمالي خرج (EGP)", "صافي التدفق (EGP)", "صافي التدفق (USD)", "صافي الأجانب (EGP)"]
    for col_idx, h in enumerate(weekly_headers, 1):
        c = ws_week.cell(row=3, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_center

    weekly_rows = get_weekly_summary(sb)
    for r_idx, r in enumerate(weekly_rows, 4):
        ws_week.cell(row=r_idx, column=1, value=f"{r['start_date']} إلى {r['end_date']}").alignment = align_center
        ws_week.cell(row=r_idx, column=2, value=r["trading_days"]).alignment = align_center
        ws_week.cell(row=r_idx, column=3, value=r["total_inflow_egp"]).number_format = '#,##0'
        ws_week.cell(row=r_idx, column=4, value=r["total_outflow_egp"]).number_format = '#,##0'
        
        c_net = ws_week.cell(row=r_idx, column=5, value=r["total_net_egp"])
        c_net.number_format = '#,##0'
        c_net.fill = green_fill if r["total_net_egp"] >= 0 else red_fill
        c_net.font = green_font if r["total_net_egp"] >= 0 else red_font

        c_net_u = ws_week.cell(row=r_idx, column=6, value=r["total_net_usd"])
        c_net_u.number_format = '$#,##0.00'
        c_net_u.fill = green_fill if r["total_net_usd"] >= 0 else red_fill
        c_net_u.font = green_font if r["total_net_usd"] >= 0 else red_font

        c_f = ws_week.cell(row=r_idx, column=7, value=r["foreign_net_egp"])
        c_f.number_format = '#,##0'
        c_f.fill = green_fill if r["foreign_net_egp"] >= 0 else red_fill
        c_f.font = green_font if r["foreign_net_egp"] >= 0 else red_font

    # 3. Monthly Sheet
    ws_month = wb.create_sheet(title="المقارنة الشهرية")
    ws_month.sheet_view.rightToLeft = True
    ws_month.merge_cells("A1:G1")
    ws_month["A1"].value = "ملخص تدفقات المؤسسات شهرياً"
    ws_month["A1"].font = title_font
    ws_month["A1"].fill = title_fill
    ws_month["A1"].alignment = align_center

    monthly_headers = ["الشهر", "أيام التداول", "إجمالي دخل (EGP)", "إجمالي خرج (EGP)", "صافي التدفق (EGP)", "صافي التدفق (USD)", "صافي الأجانب (EGP)"]
    for col_idx, h in enumerate(monthly_headers, 1):
        c = ws_month.cell(row=3, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_center

    monthly_rows = get_monthly_summary(sb)
    for r_idx, r in enumerate(monthly_rows, 4):
        ws_month.cell(row=r_idx, column=1, value=r["month_period"]).alignment = align_center
        ws_month.cell(row=r_idx, column=2, value=r["trading_days"]).alignment = align_center
        ws_month.cell(row=r_idx, column=3, value=r["total_inflow_egp"]).number_format = '#,##0'
        ws_month.cell(row=r_idx, column=4, value=r["total_outflow_egp"]).number_format = '#,##0'
        
        c_net = ws_month.cell(row=r_idx, column=5, value=r["total_net_egp"])
        c_net.number_format = '#,##0'
        c_net.fill = green_fill if r["total_net_egp"] >= 0 else red_fill
        c_net.font = green_font if r["total_net_egp"] >= 0 else red_font

        c_net_u = ws_month.cell(row=r_idx, column=6, value=r["total_net_usd"])
        c_net_u.number_format = '$#,##0.00'
        c_net_u.fill = green_fill if r["total_net_usd"] >= 0 else red_fill
        c_net_u.font = green_font if r["total_net_usd"] >= 0 else red_font

        c_f = ws_month.cell(row=r_idx, column=7, value=r["foreign_net_egp"])
        c_f.number_format = '#,##0'
        c_f.fill = green_fill if r["foreign_net_egp"] >= 0 else red_fill
        c_f.font = green_font if r["foreign_net_egp"] >= 0 else red_font

    # Adjust widths for all sheets
    for ws in [ws_daily, ws_week, ws_month]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if cell.row > 2 and len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
    wb.save(filepath)
    return filepath
